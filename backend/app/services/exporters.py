"""Excel and PDF rendering for any ReportResult."""
from __future__ import annotations

import io
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.services.reports import ReportResult

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
MONEY_FORMAT = "#,##0.00"


def to_excel(report: ReportResult) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = report.title[:31]

    sheet.append([report.title])
    sheet["A1"].font = Font(size=14, bold=True)
    sheet.append([report.description])
    sheet.append([f"Generated {datetime.now():%d %b %Y %H:%M}"])
    if report.filters:
        readable = ", ".join(f"{k}: {v}" for k, v in report.filters.items() if v)
        sheet.append([f"Filters — {readable}" if readable else "No filters applied"])
    sheet.append([])

    header_row = sheet.max_row + 1
    sheet.append([c.label for c in report.columns])
    for index in range(1, len(report.columns) + 1):
        cell = sheet.cell(row=header_row, column=index)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    for row in report.rows:
        values = []
        for column in report.columns:
            value = row.get(column.key)
            if column.kind == "money" and value not in (None, ""):
                try:
                    value = float(Decimal(str(value)))
                except Exception:
                    pass
            values.append(value)
        sheet.append(values)

    for index, column in enumerate(report.columns, start=1):
        letter = get_column_letter(index)
        longest = max(
            [len(column.label)] + [len(str(r.get(column.key, ""))) for r in report.rows[:400]]
        )
        sheet.column_dimensions[letter].width = min(max(longest + 3, 11), 52)
        if column.kind == "money":
            for row_index in range(header_row + 1, header_row + 1 + len(report.rows)):
                sheet.cell(row=row_index, column=index).number_format = MONEY_FORMAT

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    if report.totals:
        sheet.append([])
        sheet.append(["Totals — currencies are never combined"])
        sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
        for group, values in report.totals.items():
            if isinstance(values, dict):
                for key, value in values.items():
                    sheet.append([f"{group} — {key}", float(Decimal(str(value)))])
                    sheet.cell(row=sheet.max_row, column=2).number_format = MONEY_FORMAT
            else:
                sheet.append([group, values])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def to_pdf(report: ReportResult) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
        title=report.title,
    )
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle(
        "cell", parent=styles["BodyText"], fontSize=7.5, leading=9.5, spaceAfter=0
    )
    head_style = ParagraphStyle(
        "head", parent=cell_style, textColor=colors.white, fontName="Helvetica-Bold"
    )

    story = [
        Paragraph(report.title, styles["Title"]),
        Paragraph(report.description, styles["BodyText"]),
        Paragraph(
            f"Generated {datetime.now():%d %b %Y %H:%M} — {len(report.rows)} row(s)",
            styles["BodyText"],
        ),
    ]
    if report.filters:
        readable = ", ".join(f"<b>{k}</b>: {v}" for k, v in report.filters.items() if v)
        if readable:
            story.append(Paragraph(readable, styles["BodyText"]))
    story.append(Spacer(1, 6))

    data = [[Paragraph(c.label, head_style) for c in report.columns]]
    for row in report.rows[:4000]:
        data.append(
            [Paragraph(str(row.get(c.key, "") or ""), cell_style) for c in report.columns]
        )

    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D1D5DB")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1),
                 [colors.white, colors.HexColor("#F3F4F6")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    story.append(table)

    if len(report.rows) > 4000:
        story.append(Spacer(1, 6))
        story.append(
            Paragraph(
                f"Showing the first 4,000 of {len(report.rows)} rows. "
                f"Export to Excel for the complete set.",
                styles["BodyText"],
            )
        )

    if report.totals:
        story.append(Spacer(1, 10))
        story.append(Paragraph("Totals", styles["Heading3"]))
        lines = []
        for group, values in report.totals.items():
            if isinstance(values, dict):
                for key, value in values.items():
                    lines.append([group.replace("_", " "), key, str(value)])
            else:
                lines.append([group.replace("_", " "), "", str(values)])
        totals_table = Table([["Group", "Currency", "Amount"], *lines], hAlign="LEFT")
        totals_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5E7EB")),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D1D5DB")),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                ]
            )
        )
        story.append(totals_table)
        story.append(Spacer(1, 4))
        story.append(
            Paragraph(
                "<i>USD and cordoba figures are reported separately and are never added "
                "together.</i>",
                cell_style,
            )
        )

    doc.build(story)
    return buffer.getvalue()
