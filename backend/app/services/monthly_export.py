"""The Monthly Records workbook.

Three sheets — Summary (day 1 to the last day), Payments (every line), By shop.
Amounts are written in their own currency as numbers; every USD figure is a
formula reading the single rate cell on Summary, so changing the rate in Excel
recalculates the whole workbook and the arithmetic can be checked by anyone."""
from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.monthly import MonthData, month_label

FONT = "Arial"
MONEY = "#,##0.00;(#,##0.00);-"
COUNT = "0;-0;-"
RATE_FORMAT = "#,##0.0000"

TITLE = Font(name=FONT, size=14, bold=True)
BOLD = Font(name=FONT, bold=True)
BODY = Font(name=FONT)
MUTED = Font(name=FONT, italic=True, color="666666", size=9)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F2937")
USD_HEAD_FILL = PatternFill("solid", fgColor="14532D")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
TOTAL_FILL = PatternFill("solid", fgColor="E5E7EB")
WEEKEND_FILL = PatternFill("solid", fgColor="F3F4F6")
TOP = Border(top=Side(style="thin", color="000000"))

RATE_CELL = "$B$4"          # on Summary
RATE_REF = f"Summary!{RATE_CELL}"
USD_SUMMARY_COLUMNS = {6, 11, 14}


def _n(value) -> float:
    return float(value)


def _header(sheet, row: int, labels: list[str], usd_columns: set[int]) -> None:
    for index, label in enumerate(labels, start=1):
        cell = sheet.cell(row=row, column=index, value=label)
        cell.font = HEAD
        cell.fill = USD_HEAD_FILL if index in usd_columns else HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 30


def _usd(usd_cell: str, nio_cell: str, rate: str = RATE_CELL) -> str:
    # Empty, not an error, until a rate exists.
    return f'=IF({rate}>0,{usd_cell}+{nio_cell}/{rate},"")'


def _widths(sheet, widths: dict[int, float]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width


def _summary(book: Workbook, data: MonthData, generated: datetime) -> None:
    sheet = book.active
    sheet.title = "Summary"

    scope = data.shop.name if data.shop else "All shops"
    sheet["A1"] = f"Monthly record — {month_label(data.month)}"
    sheet["A1"].font = TITLE
    sheet["A2"] = f"{scope} · generated {generated:%d %b %Y %H:%M}"
    sheet["A2"].font = MUTED

    sheet["A4"] = "Rate: C$ per $1"
    sheet["A4"].font = BOLD
    rate_cell = sheet["B4"]
    rate_cell.fill = INPUT_FILL
    rate_cell.font = Font(name=FONT, bold=True, color="0000FF")
    rate_cell.number_format = RATE_FORMAT
    if data.rate:
        rate_cell.value = _n(data.rate.nio_per_usd)
        note = (f"Set in Reconcilia → Monthly Records on {data.rate.updated_at:%d %b %Y}. "
                "Change it here to recalculate every USD column in this workbook.")
    else:
        note = ("No rate was set for this month in Reconcilia, so the USD columns are empty. "
                "Type the month's rate here and they fill in.")
        sheet["C4"] = "← not set: enter the month's rate"
        sheet["C4"].font = Font(name=FONT, bold=True, color="C00000")
    rate_cell.comment = Comment(note, "Reconcilia")
    sheet["A5"] = ("USD columns = USD amount + C$ amount ÷ exchange rate. Original amounts are kept "
                   "in their own currency; bank matching never uses this rate.")
    sheet["A5"].font = MUTED

    labels = [
        "Date", "Day", "Sheets",
        "Sales $", "Sales C$", "Sales (USD)",
        "Bank $", "Bank C$", "Cash $", "Cash C$", "Total received (USD)",
        "Expenses $", "Expenses C$", "Expenses (USD)",
        "Matched", "Possible", "Unmatched", "Cash pending deposit",
    ]
    head = 7
    _header(sheet, head, labels, usd_columns=USD_SUMMARY_COLUMNS)

    first = head + 1
    for offset, (day, bucket) in enumerate(data.days):
        r = first + offset
        values = [
            day, day.strftime("%a"), bucket.sheets,
            _n(bucket.sales["USD"]), _n(bucket.sales["NIO"]), _usd(f"D{r}", f"E{r}"),
            _n(bucket.bank["USD"]), _n(bucket.bank["NIO"]),
            _n(bucket.cash["USD"]), _n(bucket.cash["NIO"]),
            f'=IF({RATE_CELL}>0,G{r}+I{r}+(H{r}+J{r})/{RATE_CELL},"")',
            _n(bucket.expenses["USD"]), _n(bucket.expenses["NIO"]), _usd(f"L{r}", f"M{r}"),
            bucket.status.get("MATCHED", 0), bucket.status.get("POSSIBLE", 0),
            bucket.status.get("UNMATCHED", 0), bucket.status.get("PENDING_DEPOSIT", 0),
        ]
        weekend = day.weekday() >= 5
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=r, column=column, value=value)
            cell.font = BODY
            if weekend:
                cell.fill = WEEKEND_FILL
            if column == 1:
                cell.number_format = "dd mmm yyyy"
            elif column == 3 or column >= 15:
                cell.number_format = COUNT
            elif column >= 4:
                cell.number_format = MONEY

    last = first + len(data.days) - 1
    total_row = last + 1
    sheet.cell(row=total_row, column=1, value="Month total")
    for column in range(3, len(labels) + 1):
        letter = get_column_letter(column)
        total = f"SUM({letter}{first}:{letter}{last})"
        if column in USD_SUMMARY_COLUMNS:
            total = f'IF({RATE_CELL}>0,{total},"")'
        sheet.cell(row=total_row, column=column, value=f"={total}")
    for column in range(1, len(labels) + 1):
        cell = sheet.cell(row=total_row, column=column)
        cell.font = BOLD
        cell.fill = TOTAL_FILL
        cell.border = TOP
        if column == 3 or column >= 15:
            cell.number_format = COUNT
        elif column >= 4:
            cell.number_format = MONEY

    sheet.freeze_panes = sheet.cell(row=first, column=3)
    _widths(sheet, {1: 13, 2: 11, 3: 8, **{c: 13 for c in range(4, 15)}, **{c: 12 for c in range(15, 19)}})
    sheet.column_dimensions["A"].width = 16   # also holds the rate label
    for column in (6, 11, 14):
        sheet.column_dimensions[get_column_letter(column)].width = 15


def _payments(book: Workbook, data: MonthData) -> None:
    sheet = book.create_sheet("Payments")
    labels = [
        "Date", "Shop", "Method", "Bank", "Currency", "Amount", "Amount (USD)",
        "Reference", "Note", "Status", "Banked on", "Bank description",
    ]
    _header(sheet, 1, labels, usd_columns={7})
    status_words = {
        "MATCHED": "Matched", "POSSIBLE": "Possible match", "UNMATCHED": "Unmatched",
        "PENDING_DEPOSIT": "Cash — pending deposit", "IGNORED": "Ignored",
    }
    for offset, line in enumerate(data.payments):
        r = 2 + offset
        values = [
            line.business_date, line.shop, line.method, line.bank, line.currency,
            _n(line.amount),
            f'=IF({RATE_REF}>0,IF(E{r}="USD",F{r},F{r}/{RATE_REF}),"")',
            line.reference, line.note, status_words.get(line.status, line.status),
            line.bank_date, line.bank_description,
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=r, column=column, value=value)
            cell.font = BODY
            if column in (1, 11) and value:
                cell.number_format = "dd mmm yyyy"
            elif column in (6, 7):
                cell.number_format = MONEY

    if data.payments:
        last = 1 + len(data.payments)
        total = last + 1
        sheet.cell(row=total, column=1, value="Total (USD)")
        sheet.cell(row=total, column=7, value=f'=IF({RATE_REF}>0,SUM(G2:G{last}),"")')
        for column in range(1, len(labels) + 1):
            cell = sheet.cell(row=total, column=column)
            cell.font = BOLD
            cell.fill = TOTAL_FILL
            cell.border = TOP
        sheet.cell(row=total, column=7).number_format = MONEY
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(labels))}{last}"
    else:
        sheet.cell(row=2, column=1, value="No payments recorded this month.").font = MUTED

    sheet.freeze_panes = "A2"
    _widths(sheet, {1: 13, 2: 22, 3: 9, 4: 10, 5: 9, 6: 13, 7: 14, 8: 16, 9: 24,
                    10: 22, 11: 13, 12: 40})


def _by_shop(book: Workbook, data: MonthData) -> None:
    sheet = book.create_sheet("By shop")
    labels = [
        "Shop", "Sheets", "Sales $", "Sales C$", "Sales (USD)",
        "Bank $", "Bank C$", "Cash $", "Cash C$", "Total received (USD)",
        "Expenses $", "Expenses C$", "Expenses (USD)",
    ]
    _header(sheet, 1, labels, usd_columns={5, 10, 13})
    for offset, (name, bucket) in enumerate(data.by_shop):
        r = 2 + offset
        values = [
            name, bucket.sheets,
            _n(bucket.sales["USD"]), _n(bucket.sales["NIO"]), _usd(f"C{r}", f"D{r}", RATE_REF),
            _n(bucket.bank["USD"]), _n(bucket.bank["NIO"]),
            _n(bucket.cash["USD"]), _n(bucket.cash["NIO"]),
            f'=IF({RATE_REF}>0,F{r}+H{r}+(G{r}+I{r})/{RATE_REF},"")',
            _n(bucket.expenses["USD"]), _n(bucket.expenses["NIO"]),
            _usd(f"K{r}", f"L{r}", RATE_REF),
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=r, column=column, value=value)
            cell.font = BODY
            if column >= 3:
                cell.number_format = MONEY

    if data.by_shop:
        first, last = 2, 1 + len(data.by_shop)
        total = last + 1
        sheet.cell(row=total, column=1, value="All shops")
        for column in range(2, len(labels) + 1):
            letter = get_column_letter(column)
            formula = f"SUM({letter}{first}:{letter}{last})"
            if column in (5, 10, 13):
                formula = f'IF({RATE_REF}>0,{formula},"")'
            sheet.cell(row=total, column=column, value=f"={formula}")
        for column in range(1, len(labels) + 1):
            cell = sheet.cell(row=total, column=column)
            cell.font = BOLD
            cell.fill = TOTAL_FILL
            cell.border = TOP
            if column >= 3:
                cell.number_format = MONEY
    else:
        sheet.cell(row=2, column=1, value="No sheets this month.").font = MUTED

    sheet.freeze_panes = "B2"
    _widths(sheet, {1: 24, 2: 8, **{c: 13 for c in range(3, 14)}})
    for column in (5, 10, 13):
        sheet.column_dimensions[get_column_letter(column)].width = 15


def _printable(sheet) -> None:
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def build_workbook(data: MonthData, generated: datetime | None = None) -> bytes:
    book = Workbook()
    _summary(book, data, generated or datetime.now())
    _payments(book, data)
    _by_shop(book, data)
    for sheet in book.worksheets:
        _printable(sheet)
    # openpyxl stores formulas without results; make Excel compute on open.
    book.calculation.fullCalcOnLoad = True
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
