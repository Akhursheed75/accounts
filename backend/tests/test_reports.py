from __future__ import annotations

from decimal import Decimal

from tests.conftest import make_sheet
from tests.test_reconciliation import add_bank_transaction, status_of


def seed_a_month(client, admin_headers, world, db):
    add_bank_transaction(
        db, world, account_label="BAC Dollars", amount="480.00",
        txn_date="2026-08-25", currency="USD",
    )
    add_bank_transaction(
        db, world, account_label="BAC Cordobas", amount="7000.00",
        txn_date="2026-08-25", currency="NIO", description="Nobody claimed this",
    )
    make_sheet(
        client, admin_headers, shop_id=world["shops"]["SHOP1"],
        business_date="2026-08-25", sales_usd="1000.00", sales_nio="20000.00",
        transfers=[
            {"bank_id": world["banks"]["BAC"], "currency_code": "USD", "amount": "480.00"},
            {"bank_id": world["banks"]["LAFISE"], "currency_code": "NIO", "amount": "9999.00"},
        ],
    )


def test_every_report_runs(client, admin_headers, world, db):
    seed_a_month(client, admin_headers, world, db)
    keys = [row["key"] for row in client.get("/api/v1/reports", headers=admin_headers).json()]
    assert len(keys) == 10
    for key in keys:
        response = client.get(
            f"/api/v1/reports/{key}",
            params={"date_from": "2026-08-01", "date_to": "2026-08-31"},
            headers=admin_headers,
        )
        assert response.status_code == 200, key
        body = response.json()
        assert body["columns"], key
        assert isinstance(body["rows"], list), key


def test_an_unknown_report_says_which_ones_exist(client, admin_headers):
    response = client.get("/api/v1/reports/not-a-report", headers=admin_headers)
    assert response.status_code == 404
    assert "daily-accounting" in response.json()["error"]["message"]


def test_reconciliation_summary_splits_by_currency(client, admin_headers, world, db):
    seed_a_month(client, admin_headers, world, db)
    body = client.get(
        "/api/v1/reports/reconciliation-summary",
        params={"date_from": "2026-08-01", "date_to": "2026-08-31"},
        headers=admin_headers,
    ).json()
    assert body["totals"]["USD"]["matched"] == "480.00"
    assert body["totals"]["NIO"]["unmatched"] == "9999.00"
    # There is no key that would hold a mixed-currency figure.
    assert set(body["totals"]) <= {"USD", "NIO"}


def test_unmatched_report_lists_both_directions(client, admin_headers, world, db):
    seed_a_month(client, admin_headers, world, db)
    body = client.get(
        "/api/v1/reports/unmatched-transactions",
        params={"date_from": "2026-08-01", "date_to": "2026-08-31"},
        headers=admin_headers,
    ).json()
    sides = {row["side"] for row in body["rows"]}
    assert "Shop payment not found in bank" in sides
    assert "Bank receipt not found in shop records" in sides


def test_currency_specific_reports_only_contain_that_currency(client, admin_headers, world, db):
    seed_a_month(client, admin_headers, world, db)
    usd = client.get("/api/v1/reports/usd-transactions", headers=admin_headers).json()
    assert usd["rows"] and all(row["currency"] == "USD" for row in usd["rows"])
    nio = client.get("/api/v1/reports/nio-transactions", headers=admin_headers).json()
    assert nio["rows"] and all(row["currency"] == "NIO" for row in nio["rows"])


def test_excel_export_is_a_real_workbook(client, admin_headers, world, db):
    from io import BytesIO

    from openpyxl import load_workbook

    seed_a_month(client, admin_headers, world, db)
    response = client.get(
        "/api/v1/reports/daily-accounting",
        params={"format": "xlsx", "date_from": "2026-08-01", "date_to": "2026-08-31"},
        headers=admin_headers,
    )
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    values = [
        str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None
    ]
    assert any("Daily accounting" in value for value in values)
    assert any("1000" in value for value in values)


def test_pdf_export_is_a_real_pdf(client, admin_headers, world, db):
    seed_a_month(client, admin_headers, world, db)
    response = client.get(
        "/api/v1/reports/reconciliation-summary",
        params={"format": "pdf"}, headers=admin_headers,
    )
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF-")
    assert len(response.content) > 1000


def test_exports_are_audited(client, admin_headers, world, db):
    seed_a_month(client, admin_headers, world, db)
    client.get("/api/v1/reports/shop-sales", params={"format": "xlsx"}, headers=admin_headers)
    logs = client.get(
        "/api/v1/audit-logs", params={"action": "EXPORT_REPORT"}, headers=admin_headers
    ).json()["items"]
    assert logs and "XLSX" in logs[0]["summary"]


def test_dashboard_never_combines_currencies(client, admin_headers, world, db):
    seed_a_month(client, admin_headers, world, db)
    body = client.get(
        "/api/v1/dashboard",
        params={"date_from": "2026-08-01", "date_to": "2026-08-31"},
        headers=admin_headers,
    ).json()
    assert body["sales"]["USD"] == "1000.00"
    assert body["sales"]["NIO"] == "20000.00"
    assert set(body["shop_transfers"]) == {"USD", "NIO"}
    assert body["reconciliation"]["counts"]["MATCHED"] == 1
    assert body["reconciliation"]["counts"]["UNMATCHED"] == 1
    assert body["reconciliation"]["amounts"]["MATCHED"]["USD"] == "480.00"
    assert body["reconciliation"]["amounts"]["UNMATCHED"]["NIO"] == "9999.00"


def test_a_shop_user_report_only_covers_their_shop(client, admin_headers, shop_headers, world, db):
    seed_a_month(client, admin_headers, world, db)
    make_sheet(
        client, admin_headers, shop_id=world["shops"]["SHOP2"],
        business_date="2026-08-26", sales_usd="5555.00", transfers=[],
    )
    body = client.get("/api/v1/reports/shop-sales", headers=shop_headers).json()
    shops = {row["shop"] for row in body["rows"]}
    assert shops == {"Shop 1"}
