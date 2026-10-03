"""Monthly Records: the current month and the three before it, day by day, in USD.

Dates are taken relative to today so the window rule keeps being tested as the
calendar moves, rather than passing only in the month it was written."""
from __future__ import annotations

import calendar
from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app.services.monthly import add_months, available_months, month_key
from tests.conftest import make_sheet

PREVIOUS = available_months()[1]                 # always a complete month in the past
KEY = month_key(PREVIOUS)
DAY3 = PREVIOUS.replace(day=3).isoformat()


def sheet_on_day_three(client, headers, world, shop="SHOP1"):
    return make_sheet(
        client, headers, shop_id=world["shops"][shop], business_date=DAY3,
        sales_usd="100.00", sales_nio="3660.00",
        transfers=[
            {"bank_id": world["banks"]["BAC"], "currency_code": "USD", "amount": "50.00"},
            {"payment_method": "CASH", "currency_code": "NIO", "amount": "1830.00"},
        ],
    ).json()


def test_lists_the_current_month_and_the_three_before_it(client, admin_headers):
    body = client.get("/api/v1/monthly", headers=admin_headers).json()
    months = [m["month"] for m in body["months"]]
    today = date.today().replace(day=1)
    assert months == [month_key(add_months(today, -i)) for i in range(4)]
    assert body["months"][0]["is_current"] is True
    assert body["previous_months"] == 3


def test_a_month_older_than_the_window_is_refused_with_an_explanation(client, admin_headers):
    too_old = month_key(add_months(date.today().replace(day=1), -4))
    response = client.get(f"/api/v1/monthly/{too_old}", headers=admin_headers)
    assert response.status_code == 422
    assert "not deleted" in response.text


def test_nonsense_month_is_refused(client, admin_headers):
    assert client.get("/api/v1/monthly/2026-13", headers=admin_headers).status_code == 422
    assert client.get("/api/v1/monthly/october", headers=admin_headers).status_code == 422


def test_every_day_of_the_month_is_listed_with_original_currencies(
    client, admin_headers, world
):
    sheet_on_day_three(client, admin_headers, world)
    body = client.get(f"/api/v1/monthly/{KEY}", headers=admin_headers).json()

    assert len(body["days"]) == calendar.monthrange(PREVIOUS.year, PREVIOUS.month)[1]
    assert [d["day"] for d in body["days"]][:3] == [1, 2, 3]
    day3 = body["days"][2]
    assert day3["sheets"] == 1
    assert day3["sales_usd"] == "100.00" and day3["sales_nio"] == "3660.00"
    assert day3["bank_usd"] == "50.00" and day3["cash_nio"] == "1830.00"
    assert day3["pending_cash"] == 1
    assert body["days"][0]["sheets"] == 0
    # No rate yet: the USD view is empty, never guessed.
    assert body["rate"] is None
    assert day3["sales_in_usd"] is None and body["totals"]["received_in_usd"] is None


def test_with_a_rate_everything_is_also_shown_in_usd(client, admin_headers, world):
    sheet_on_day_three(client, admin_headers, world)
    saved = client.put(
        f"/api/v1/monthly/{KEY}/rate", json={"nio_per_usd": "36.6"}, headers=admin_headers
    )
    assert saved.status_code == 200, saved.text

    body = client.get(f"/api/v1/monthly/{KEY}", headers=admin_headers).json()
    assert body["rate"] == "36.6000"
    day3 = body["days"][2]
    assert day3["sales_in_usd"] == "200.00"        # 100 + 3660 / 36.6
    assert day3["bank_in_usd"] == "50.00"
    assert day3["cash_in_usd"] == "50.00"          # 1830 / 36.6
    assert day3["received_in_usd"] == "100.00"
    assert body["totals"]["received_in_usd"] == "100.00"
    # The originals are untouched by the rate.
    assert body["totals"]["cash_nio"] == "1830.00"


def test_changing_the_rate_replaces_it_and_is_audited(client, admin_headers, world):
    client.put(f"/api/v1/monthly/{KEY}/rate", json={"nio_per_usd": "36.6"}, headers=admin_headers)
    client.put(f"/api/v1/monthly/{KEY}/rate", json={"nio_per_usd": "36.7"}, headers=admin_headers)
    months = client.get("/api/v1/monthly", headers=admin_headers).json()["months"]
    assert next(m for m in months if m["month"] == KEY)["rate"] == "36.7000"
    log = client.get("/api/v1/audit-logs", params={"action": "SET_EXCHANGE_RATE"},
                     headers=admin_headers).json()
    summaries = [row["summary"] for row in log["items"]]
    assert any("was 36.6000" in s for s in summaries)


def test_implausible_rates_are_refused(client, admin_headers):
    for bad in ("0", "0.5", "3662", "-36"):
        response = client.put(
            f"/api/v1/monthly/{KEY}/rate", json={"nio_per_usd": bad}, headers=admin_headers
        )
        assert response.status_code == 422, bad


def test_only_an_administrator_sets_the_rate(client, shop_headers):
    response = client.put(
        f"/api/v1/monthly/{KEY}/rate", json={"nio_per_usd": "36.6"}, headers=shop_headers
    )
    assert response.status_code == 403


def test_a_shop_user_sees_only_their_own_shop(client, admin_headers, shop_headers, world):
    sheet_on_day_three(client, admin_headers, world, shop="SHOP1")
    sheet_on_day_three(client, admin_headers, world, shop="SHOP2")
    mine = client.get(f"/api/v1/monthly/{KEY}", headers=shop_headers).json()
    everyone = client.get(f"/api/v1/monthly/{KEY}", headers=admin_headers).json()
    assert mine["days"][2]["sheets"] == 1
    assert everyone["days"][2]["sheets"] == 2
    assert client.get(
        f"/api/v1/monthly/{KEY}", params={"shop_id": world["shops"]["SHOP2"]},
        headers=shop_headers,
    ).status_code == 403


def test_a_shop_user_cannot_download_the_workbook(client, shop_headers):
    assert client.get(f"/api/v1/monthly/{KEY}/export", headers=shop_headers).status_code == 403


def test_the_rate_for_any_day_is_available_to_the_daily_sheet(client, admin_headers):
    client.put(f"/api/v1/monthly/{KEY}/rate", json={"nio_per_usd": "36.6"}, headers=admin_headers)
    body = client.get("/api/v1/monthly/rate", params={"on": DAY3}, headers=admin_headers).json()
    assert body == {"month": KEY, "nio_per_usd": "36.6000"}
    old = add_months(PREVIOUS, -12).isoformat()
    assert client.get("/api/v1/monthly/rate", params={"on": old},
                      headers=admin_headers).json()["nio_per_usd"] is None


def test_the_workbook_has_a_summary_payments_and_by_shop_sheet(client, admin_headers, world):
    sheet_on_day_three(client, admin_headers, world)
    client.put(f"/api/v1/monthly/{KEY}/rate", json={"nio_per_usd": "36.6"}, headers=admin_headers)
    response = client.get(f"/api/v1/monthly/{KEY}/export", headers=admin_headers)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert f"reconcilia-{KEY}" in response.headers["content-disposition"]

    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Summary", "Payments", "By shop"]
    summary = book["Summary"]
    assert summary["B4"].value == 36.6
    days = calendar.monthrange(PREVIOUS.year, PREVIOUS.month)[1]
    # Header on row 7, one row per day, then the total.
    assert summary.cell(row=8 + days, column=1).value == "Month total"
    day3_row = 10
    assert summary.cell(row=day3_row, column=4).value == 100.0
    assert summary.cell(row=day3_row, column=6).value.startswith("=IF($B$4>0")

    payments = book["Payments"]
    methods = {payments.cell(row=r, column=3).value for r in (2, 3)}
    assert methods == {"Bank", "Cash"}
